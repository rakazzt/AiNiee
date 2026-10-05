"""Regression tests for the XLSX writer fixes reviewed before merging.

Covers the P0 security regression found in review: the rewrite to template-based
write-back dropped the guard that stopped a translation beginning with "=" from
being stored as a live formula (CWE-1236). A translation's text comes from a model
whose input can be a third-party file, so it is untrusted.

Also pins the two properties the review asked for: the user's source workbook is
never modified, and the rest of the workbook survives the round-trip.
"""
import hashlib
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl

from ModuleFolders.Domain.FileOutputer.XlsxWriter import XlsxWriter
from ModuleFolders.Service.Cache.CacheFile import CacheFile
from ModuleFolders.Service.Cache.CacheItem import CacheItem, TranslationStatus

DATACOLS = 3


def _content_cell(sheet, row_index, col_index):
    """The writer offsets by +2 rows (header + 1-based) and +1 column."""
    return sheet.cell(row=row_index + 2, column=col_index + 1)


class XlsxWriterInjectionTests(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.source = self.tmp / "source.xlsx"
        self.output = self.tmp / "out.xlsx"
        self.writer = XlsxWriter.__new__(XlsxWriter)  # on_write_translated uses no instance state

    def tearDown(self):
        self._tmp.cleanup()

    def _make_source(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet"
        ws.append(["h1", "h2", "h3"])
        ws.append(["plain text", 42, "=SUM(B2:B2)"])
        ws.append(["second row", 3.5, "tail"])
        wb.create_sheet("Other")
        wb["Other"]["A1"] = "untouched sheet"
        wb.save(self.source)
        return wb

    def _cache(self, items):
        cache = CacheFile(items=[])
        for source_text, translated, row, col in items:
            item = CacheItem(source_text=source_text, translated_text=translated,
                             extra={"row": row, "col": col})
            item.translation_status = TranslationStatus.TRANSLATED
            cache.items.append(item)
        return cache

    def _run(self, items):
        self._make_source()
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.writer.on_write_translated(
            self.output, self._cache(items), None, source_file_path=self.source
        )
        after = hashlib.sha256(self.source.read_bytes()).hexdigest()
        return before, after, openpyxl.load_workbook(self.output)

    def test_translation_starting_with_equals_is_stored_as_text(self):
        """P0: '='-prefixed translations must not become live formulas."""
        _, _, wb = self._run([("plain text", "=1+1", 0, 0)])
        cell = _content_cell(wb["Sheet"], 0, 0)
        self.assertNotEqual(cell.data_type, "f", "translation was stored as an executable formula")
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, "=1+1", "text must be preserved verbatim")

    def test_dde_style_payload_is_stored_as_text(self):
        _, _, wb = self._run([("plain text", "=cmd|'/c calc'!A0", 0, 0)])
        cell = _content_cell(wb["Sheet"], 0, 0)
        self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(cell.value, "=cmd|'/c calc'!A0")

    def test_source_workbook_is_never_modified(self):
        before, after, _ = self._run([("plain text", "译文", 0, 0)])
        self.assertEqual(before, after, "the user's source workbook must not be rewritten")

    def test_user_formula_cell_is_not_overwritten(self):
        """A cache item pointing at a formula cell must not clobber the user's formula."""
        _, _, wb = self._run([("plain text", "译文", 0, 0), ("formula", "覆盖", 0, 2)])
        formula_cell = _content_cell(wb["Sheet"], 0, 2)
        self.assertEqual(formula_cell.value, "=SUM(B2:B2)", "user formula was overwritten")
        self.assertEqual(formula_cell.data_type, "f")

    def test_other_sheets_and_headers_survive(self):
        _, _, wb = self._run([("plain text", "译文", 0, 0)])
        self.assertIn("Other", wb.sheetnames)
        self.assertEqual(wb["Other"]["A1"].value, "untouched sheet")
        self.assertEqual(wb["Sheet"].cell(row=1, column=1).value, "h1")

    def test_translated_string_cell_is_written(self):
        _, _, wb = self._run([("plain text", "译文", 0, 0)])
        self.assertEqual(_content_cell(wb["Sheet"], 0, 0).value, "译文")

    def test_untranslated_items_are_left_alone(self):
        self._make_source()
        item = CacheItem(source_text="second row", translated_text="不该写入",
                         extra={"row": 1, "col": 0})
        item.translation_status = TranslationStatus.UNTRANSLATED
        self.writer.on_write_translated(
            self.output, CacheFile(items=[item]), None, source_file_path=self.source
        )
        wb = openpyxl.load_workbook(self.output)
        self.assertEqual(_content_cell(wb["Sheet"], 1, 0).value, "second row")


if __name__ == "__main__":
    unittest.main()
